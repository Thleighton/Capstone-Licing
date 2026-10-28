from io import BytesIO
from unittest.mock import AsyncMock, Mock

from openpyxl import load_workbook

from app.api import preseleccion
from app.database import get_db
from app.main import app
from app.services import filtradas
from app.services.exportacion_oportunidades import crear_excel


def test_excel_tres_hojas_enlaces_y_nombres_como_texto():
    contenido = crear_excel([
        {"codigo_externo": "123-1-LE26", "nombre_organismo": '=HYPERLINK("https://malicioso.test")', "estado_seleccion": "preseleccionada"},
        {"codigo_externo": "123-2-LE26", "nombre_unidad": "Educación\x00", "estado_seleccion": "revision_manual"},
        {"codigo_externo": "123-3-LE26", "estado_seleccion": "descartada"},
    ])
    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Posibles oportunidades", "Pendientes de revisión", "Descartadas"]
    for hoja in libro:
        assert hoja.max_column == 3 and hoja.max_row == 2
        assert list(hoja.values)[0] == ("ID licitación", "Organismo solicitante", "Link directo")
        assert hoja["C2"].hyperlink.target.endswith(hoja["A2"].value)
        assert hoja.freeze_panes == "A2"
    assert libro.worksheets[0]["B2"].data_type == "s"
    assert libro.worksheets[1]["B2"].value == "Educación"


def preparar(monkeypatch, cantidad=103):
    entradas = {str(i): {"id": i, "raw_id": i, "codigo_externo": f"123-{i}-LE26",
                         "nombre_organismo": "Universidad de Chile"} for i in range(1, cantidad + 1)}
    db = Mock()
    db.scalars.return_value.__iter__ = lambda _: iter(range(1, cantidad + 1))
    app.dependency_overrides[get_db] = lambda: db
    monkeypatch.setattr(preseleccion, "cargar_entradas", lambda db, ids: entradas)
    guardadas = [{**d, "fuente_raw_id": d["raw_id"], "hash_fuente": preseleccion.firma_fuente(d),
                  "estado_seleccion": ["preseleccionada", "revision_manual", "descartada"][int(i) % 3]}
                 for i, d in entradas.items()]
    return entradas, guardadas


def test_exportacion_todas_las_paginas_y_evaluaciones_obsoletas(client_usuario, monkeypatch):
    entradas, guardadas = preparar(monkeypatch)
    guardadas[0]["hash_fuente"] = "obsoleta"
    guardadas.pop(1)  # Una licitación sin evaluar también debe quedar en revisión.
    async def operar(operacion, datos):
        assert operacion == "listar"
        return {"filas": guardadas[datos["offset"]:datos["offset"] + datos["limite"]]}
    servicio = AsyncMock(side_effect=operar)
    monkeypatch.setattr(filtradas, "operar", servicio)
    r = client_usuario.get("/api/preseleccion/exportar.xlsx")
    assert r.status_code == 200, r.text[:200] if r.status_code != 200 else ""
    assert r.headers["cache-control"] == "no-store"
    assert "attachment" in r.headers["content-disposition"]
    libro = load_workbook(BytesIO(r.content))
    assert sum(h.max_row - 1 for h in libro) == len(entradas)
    revisiones = {r[0] for r in libro.worksheets[1].iter_rows(min_row=2, values_only=True)}
    assert {"123-1-LE26", "123-2-LE26"} <= revisiones
    assert servicio.await_count == 2


def test_excel_vacio_conserva_las_tres_pestanas(client_usuario, monkeypatch):
    preparar(monkeypatch, 0)
    servicio = AsyncMock()
    monkeypatch.setattr(filtradas, "operar", servicio)
    r = client_usuario.get("/api/preseleccion/exportar.xlsx")
    assert r.status_code == 200
    assert all(h.max_row == 1 for h in load_workbook(BytesIO(r.content)))
    servicio.assert_not_awaited()


def test_fallo_no_entrega_excel_incompleto(client_usuario, monkeypatch):
    preparar(monkeypatch)
    monkeypatch.setattr(filtradas, "operar", AsyncMock(side_effect=filtradas.FiltradasNoDisponible("No disponible")))
    assert client_usuario.get("/api/preseleccion/exportar.xlsx").status_code == 503


def test_exportacion_exige_sesion(client):
    assert client.get("/api/preseleccion/exportar.xlsx").status_code == 401
