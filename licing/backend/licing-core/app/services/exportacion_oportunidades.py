"""Excel de licitaciones: solo referencia, comprador y enlace público."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment

from .enlaces import ficha_oficial


HOJAS = {
    "preseleccionada": "Posibles oportunidades",
    "revision_manual": "Pendientes de revisión",
    "descartada": "Descartadas",
}


def crear_excel(filas):
    libro = Workbook()
    libro.remove(libro.active)
    filas = sorted(filas, key=lambda f: f["codigo_externo"])
    for estado, nombre in HOJAS.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(["ID licitación", "Organismo solicitante", "Link directo"])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="087D89")
        for fila in filas:
            if fila["estado_seleccion"] != estado:
                continue
            codigo = fila["codigo_externo"]
            valores = [codigo, fila.get("nombre_organismo") or fila.get("nombre_unidad") or "No informado",
                       ficha_oficial(codigo) or ""]
            hoja.append([ILLEGAL_CHARACTERS_RE.sub("", str(v)) for v in valores])
            # Los nombres de la fuente son texto, nunca fórmulas de Excel.
            for celda in hoja[hoja.max_row]:
                celda.data_type = "s"
                celda.alignment = Alignment(vertical="top", wrap_text=True)
            enlace = hoja.cell(hoja.max_row, 3)
            if enlace.value:
                enlace.hyperlink = enlace.value
                enlace.style = "Hyperlink"
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for columna, ancho in (("A", 24), ("B", 65), ("C", 95)):
            hoja.column_dimensions[columna].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
